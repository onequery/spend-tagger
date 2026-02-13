#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def main() -> None:
    src = Path("통합지출내역.xlsx")
    out = Path("월별_총지출_요약.xlsx")

    if not src.exists():
        raise FileNotFoundError(f"입력 파일이 없습니다: {src}")

    df = pd.read_excel(src)
    required = ["지출날짜", "지출금액"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"필수 컬럼 누락: {missing}")

    df["지출날짜"] = pd.to_datetime(df["지출날짜"], errors="coerce")
    df["지출금액"] = pd.to_numeric(df["지출금액"], errors="coerce")
    df = df[df["지출날짜"].notna() & df["지출금액"].notna() & (df["지출금액"] > 0)].copy()

    df["연월"] = df["지출날짜"].dt.to_period("M")
    monthly = (
        df.groupby("연월", as_index=False)
        .agg(
            총_지출금액=("지출금액", "sum"),
            거래건수=("지출금액", "size"),
        )
        .sort_values("연월")
    )
    monthly["건당_평균지출"] = monthly["총_지출금액"] / monthly["거래건수"]
    monthly["연월"] = monthly["연월"].astype(str)

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        monthly.to_excel(writer, index=False, sheet_name="월별총지출")

    # 간단한 표시 형식 정리
    wb = load_workbook(out)
    ws = wb["월별총지출"]
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 16
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
        row[3].number_format = "#,##0"
    wb.save(out)

    print(f"saved: {out} rows={len(monthly)}")
    print(monthly.to_string(index=False))


if __name__ == "__main__":
    main()
