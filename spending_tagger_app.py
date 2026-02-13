#!/usr/bin/env python3
from __future__ import annotations

import json
import os
import re
import urllib.error
import urllib.request
import warnings
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import zipfile
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
import xml.etree.ElementTree as ET

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib import font_manager, rcParams
from openpyxl import load_workbook

from merge_transactions import _parse_biff_xls, _parse_biff_xls_raw, _read_strict_xlsx

warnings.filterwarnings("ignore", message=r"Glyph .* missing from font")


EXCLUDED_CATEGORIES = {"제외", "금융/투자"}
STANDARD_COLUMNS = ["지출날짜", "사용처", "지출금액", "카테고리", "세부 카테고리", "거래구분", "출처파일"]
MEMORY_PATH = Path("tagging_memory.csv")
ONTOLOGY_PATH = Path("ontology.json")
MAPPING_PROFILES_PATH = Path("column_mapping_profiles.json")

TARGET_COLUMNS = ["지출날짜", "사용처", "지출금액", "카테고리", "세부 카테고리", "거래구분"]
REQUIRED_TARGET_COLUMNS = ["지출날짜", "사용처", "지출금액"]

DEFAULT_ONTOLOGY: dict[str, list[str]] = {
    "주거": ["원리금", "관리비", "전기/가스/수도", "인터넷", "가구/가전구입", "주택유지보수", "미분류"],
    "건강": ["보험", "병원/약국", "알레르기 병원/약국", "피부과 병원/약국", "사고/외상 치료", "치과", "미분류"],
    "차량": ["유류비", "정비/소모품(타이어,엔진오일 등)", "보험료", "세차", "주차비", "통행료", "미분류"],
    "식비": ["외식", "장보기(요리용 식재료)", "배달", "카페/디저트", "미분류"],
    "사회관계": ["선물", "축의금/부의금", "기념일/이벤트", "모임", "미분류"],
    "사회보험": ["국민연금", "건강보험", "고용보험", "산재보험", "미분류"],
    "쇼핑": ["생활용품", "의류", "전자기기", "미분류"],
    "생활비": ["통신비", "인터넷", "정기구독", "미분류"],
    "라이프이벤트": ["스튜디오/촬영", "결혼/약혼", "이사", "미분류"],
    "자기관리": ["이발", "뷰티", "운동", "미분류"],
    "여행/여가": ["관광/체험", "숙박", "교통", "미분류"],
    "교통": ["대중교통", "택시", "철도/항공", "미분류"],
    "금융/투자": ["CMA", "저축", "투자", "이체", "미분류"],
    "제외": ["급여 보상", "타행반영", "분석제외", "카드결제(월납)", "미분류"],
}

BANK_SOURCE_KEYWORDS = ("은행", "bank")
CARD_SETTLEMENT_TERMS = (
    "카드결제",
    "카드대금",
    "카드이용대금",
    "신용카드",
    "카드청구",
    "카드납부",
)
CARD_COMPANY_TERMS = (
    "현대카드",
    "신한카드",
    "삼성카드",
    "롯데카드",
    "하나카드",
    "우리카드",
    "국민카드",
    "kb국민카드",
    "농협카드",
    "nh카드",
    "비씨카드",
    "bc카드",
)
MANUAL_REVIEW_TRANSFER_TERMS = ("타행송금", "타행이체")


@dataclass
class TagSuggestion:
    category: str
    subcategory: str
    confidence: float
    source: str
    reason: str = ""


@dataclass
class KnowledgeBase:
    merchant_map: dict[str, dict[tuple[str, str], float]]
    keyword_map: dict[str, dict[tuple[str, str], float]]
    subcat_by_category: dict[str, list[str]]


def get_default_ontology() -> dict[str, list[str]]:
    return {cat: list(subs) for cat, subs in DEFAULT_ONTOLOGY.items()}


def normalize_ontology(raw: dict) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    if not isinstance(raw, dict):
        return get_default_ontology()

    for cat, subs in raw.items():
        category = str(cat).strip()
        if not category:
            continue

        cleaned: list[str] = []
        if isinstance(subs, list):
            values = subs
        elif subs is None:
            values = []
        else:
            values = [subs]

        for sub in values:
            sub_name = str(sub).strip()
            if sub_name and sub_name not in cleaned:
                cleaned.append(sub_name)

        if not cleaned:
            cleaned = ["미분류"]
        out[category] = cleaned

    if not out:
        return get_default_ontology()
    return out


def load_ontology() -> dict[str, list[str]]:
    if not ONTOLOGY_PATH.exists():
        return get_default_ontology()
    try:
        data = json.loads(ONTOLOGY_PATH.read_text(encoding="utf-8"))
        return normalize_ontology(data)
    except Exception:
        return get_default_ontology()


def save_ontology(ontology: dict[str, list[str]]) -> None:
    cleaned = normalize_ontology(ontology)
    ONTOLOGY_PATH.write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")


def get_column_aliases() -> dict[str, list[str]]:
    return {
        "지출날짜": ["지출날짜", "날짜", "거래일시", "거래일", "사용일시", "일시", "승인일", "승인일자", "거래일자"],
        "사용처": ["사용처", "거래처", "가맹점명", "가맹점", "적요", "내역", "상호", "상호명"],
        "지출금액": ["지출금액", "금액", "거래금액", "사용금액", "출금금액", "출금액", "결제금액", "승인금액", "이용금액"],
        "카테고리": ["카테고리", "분류", "대분류"],
        "세부 카테고리": ["세부 카테고리", "세부카테고리", "소분류"],
        "거래구분": ["거래구분", "구분", "이용구분", "승인구분"],
    }


def infer_alias_mapping(columns: list[str]) -> dict[str, str]:
    aliases = get_column_aliases()
    cols = [str(c).strip() for c in columns]
    mapping: dict[str, str] = {}
    for target, candidates in aliases.items():
        for c in cols:
            if c in candidates:
                mapping[target] = c
                break
    return mapping


def _normalize_profile(raw: dict) -> dict | None:
    if not isinstance(raw, dict):
        return None
    profile_name = str(raw.get("profile_name", "")).strip()
    source_keyword = str(raw.get("source_keyword", "")).strip()
    try:
        header_row = int(raw.get("header_row", 0))
    except Exception:
        header_row = 0
    column_map_raw = raw.get("column_map", {})
    if not isinstance(column_map_raw, dict):
        column_map_raw = {}

    column_map: dict[str, str] = {}
    for target in TARGET_COLUMNS:
        src = str(column_map_raw.get(target, "")).strip()
        if src and src != "(없음)":
            column_map[target] = src

    if not profile_name:
        profile_name = f"profile_{abs(hash((source_keyword, header_row))) % 10_000_000}"
    return {
        "profile_name": profile_name,
        "source_keyword": source_keyword,
        "header_row": max(0, header_row),
        "column_map": column_map,
    }


def _resolve_column_name(requested: str, candidates: list[str]) -> str:
    wanted = str(requested).strip()
    if not wanted:
        return ""
    if wanted in candidates:
        return wanted

    wanted_norm = _normalize_text(wanted).replace(" ", "")
    for cand in candidates:
        if _normalize_text(cand).replace(" ", "") == wanted_norm:
            return cand
    return ""


def validate_mapping_profile(raw_df: pd.DataFrame, profile: dict) -> dict | None:
    norm = _normalize_profile(profile)
    if norm is None or raw_df is None or raw_df.empty:
        return None

    try:
        shaped = build_dataframe_from_raw(raw_df, int(norm.get("header_row", 0)))
    except Exception:
        return None
    if shaped.empty:
        return None

    shaped_cols = [str(c).strip() for c in shaped.columns]
    resolved_map: dict[str, str] = {}
    for target, src in norm.get("column_map", {}).items():
        resolved = _resolve_column_name(str(src), shaped_cols)
        if resolved:
            resolved_map[target] = resolved

    norm["column_map"] = resolved_map
    for target in REQUIRED_TARGET_COLUMNS:
        src = str(resolved_map.get(target, "")).strip()
        if not src:
            return None
        if src not in shaped_cols:
            return None
    return norm


def load_mapping_profiles() -> list[dict]:
    if not MAPPING_PROFILES_PATH.exists():
        return []
    try:
        raw = json.loads(MAPPING_PROFILES_PATH.read_text(encoding="utf-8"))
    except Exception:
        return []
    if not isinstance(raw, list):
        return []
    out: list[dict] = []
    for item in raw:
        norm = _normalize_profile(item)
        if norm is not None:
            out.append(norm)
    return out


def save_mapping_profiles(profiles: list[dict]) -> None:
    cleaned: list[dict] = []
    seen = set()
    for p in profiles:
        norm = _normalize_profile(p)
        if norm is None:
            continue
        key = norm["profile_name"]
        if key in seen:
            continue
        seen.add(key)
        cleaned.append(norm)
    MAPPING_PROFILES_PATH.write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")


def set_korean_font() -> None:
    preferred = ["AppleGothic", "Malgun Gothic", "NanumGothic"]
    available = {f.name for f in font_manager.fontManager.ttflist}
    for name in preferred:
        if name in available:
            rcParams["font.family"] = name
            break
    rcParams["axes.unicode_minus"] = False


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("_")


def tokenize(text: str) -> list[str]:
    return [tok for tok in re.findall(r"[A-Za-z0-9가-힣]+", text.lower()) if len(tok) >= 2]


def _normalize_text(text: str) -> str:
    return unicodedata.normalize("NFC", str(text)).strip().lower()


def is_bank_source(source_file: str) -> bool:
    source = _normalize_text(source_file)
    return any(k in source for k in BANK_SOURCE_KEYWORDS)


def is_monthly_card_settlement(merchant: str, source_file: str) -> bool:
    if not is_bank_source(source_file):
        return False

    raw = _normalize_text(merchant)
    compact = raw.replace(" ", "")
    if not compact:
        return False

    if any(term in compact for term in CARD_SETTLEMENT_TERMS):
        return True

    if any(term in compact for term in CARD_COMPANY_TERMS):
        return True

    # 예: "카드사 자동이체", "카드 출금" 등
    if "카드" in compact and any(k in compact for k in ("이체", "출금", "자동", "납부", "결제", "청구")):
        return True
    return False


def requires_manual_classification(merchant: str, category: str, subcategory: str, transaction_type: str = "") -> bool:
    cat = str(category).strip()
    sub = str(subcategory).strip()
    if cat in EXCLUDED_CATEGORIES:
        return False

    compact_merchant = _normalize_text(merchant).replace(" ", "")
    compact_tx_type = _normalize_text(transaction_type).replace(" ", "")
    if any(term in compact_tx_type for term in MANUAL_REVIEW_TRANSFER_TERMS):
        return True
    if any(term in compact_merchant for term in MANUAL_REVIEW_TRANSFER_TERMS):
        return True

    if cat == "미분류" or sub == "미분류":
        return True
    return False


def _make_unique_headers(headers: list[str]) -> list[str]:
    counts: dict[str, int] = {}
    out: list[str] = []
    for i, h in enumerate(headers, start=1):
        base = str(h).strip() or f"컬럼{i}"
        n = counts.get(base, 0) + 1
        counts[base] = n
        out.append(base if n == 1 else f"{base}#{n}")
    return out


def guess_header_row(raw_df: pd.DataFrame) -> int:
    if raw_df is None or raw_df.empty:
        return 0

    hints = {
        "날짜",
        "거래일시",
        "거래일",
        "거래일자",
        "승인일",
        "승인일자",
        "지출날짜",
        "사용처",
        "거래처",
        "가맹점명",
        "가맹점",
        "상호",
        "상호명",
        "금액",
        "거래금액",
        "사용금액",
        "지출금액",
        "승인금액",
    }

    scan_limit = min(len(raw_df), 40)
    for i in range(scan_limit):
        vals = {str(v).strip() for v in raw_df.iloc[i].tolist() if str(v).strip()}
        if len(vals & hints) >= 2:
            return i

    best_idx = 0
    best_score = -1
    for i in range(min(len(raw_df), 20)):
        row_vals = [str(v).strip() for v in raw_df.iloc[i].tolist()]
        non_empty = [v for v in row_vals if v]
        unique_count = len(set(non_empty))
        score = len(non_empty) + unique_count
        if score > best_score:
            best_score = score
            best_idx = i
    return best_idx


def build_dataframe_from_raw(raw_df: pd.DataFrame, header_row: int) -> pd.DataFrame:
    if raw_df is None or raw_df.empty:
        return pd.DataFrame()
    hr = max(0, min(int(header_row), len(raw_df) - 1))
    header_vals = [str(v).strip() for v in raw_df.iloc[hr].tolist()]
    columns = _make_unique_headers(header_vals)
    body = raw_df.iloc[hr + 1 :].copy().reset_index(drop=True)
    body.columns = columns
    return body


def dataframe_to_raw_like(df: pd.DataFrame) -> pd.DataFrame:
    if df is None:
        return pd.DataFrame()
    header = pd.DataFrame([list(df.columns)])
    body = df.copy().reset_index(drop=True)
    body.columns = range(body.shape[1])
    header.columns = range(header.shape[1])
    return pd.concat([header, body], ignore_index=True)


def _read_strict_xlsx_raw(path: Path) -> pd.DataFrame:
    with zipfile.ZipFile(path) as zf:
        wb_xml = ET.fromstring(zf.read("xl/workbook.xml"))
        ns_main = {"m": wb_xml.tag.split("}")[0].strip("{")}
        rel_xml = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rel_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        rel_map = {rel.attrib.get("Id"): rel.attrib.get("Target", "") for rel in rel_xml.findall("r:Relationship", rel_ns)}

        sheet_el = wb_xml.find("m:sheets/m:sheet", ns_main)
        if sheet_el is None:
            return pd.DataFrame()

        rid = sheet_el.attrib.get("{http://purl.oclc.org/ooxml/officeDocument/relationships}id") or sheet_el.attrib.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        if not rid or rid not in rel_map:
            return pd.DataFrame()

        sheet_path = "xl/" + rel_map[rid]
        if sheet_path.startswith("xl//"):
            sheet_path = "xl/" + sheet_path[4:]

        shared: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            sst_xml = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            sst_ns = {"s": sst_xml.tag.split("}")[0].strip("{")}
            for si in sst_xml.findall("s:si", sst_ns):
                chunks = []
                for t in si.findall(".//s:t", sst_ns):
                    chunks.append(t.text or "")
                shared.append("".join(chunks))

        sheet_xml = ET.fromstring(zf.read(sheet_path))
        s_ns = {"s": sheet_xml.tag.split("}")[0].strip("{")}
        rows: dict[int, dict[int, object]] = {}

        def col_to_idx(col: str) -> int:
            n = 0
            for ch in col:
                if "A" <= ch <= "Z":
                    n = n * 26 + (ord(ch) - ord("A") + 1)
            return n - 1

        for row in sheet_xml.findall(".//s:sheetData/s:row", s_ns):
            r_idx = int(row.attrib.get("r", "1")) - 1
            rmap = rows.setdefault(r_idx, {})
            for c in row.findall("s:c", s_ns):
                ref = c.attrib.get("r", "")
                col = "".join(ch for ch in ref if ch.isalpha())
                c_idx = col_to_idx(col) if col else 0
                c_type = c.attrib.get("t")
                v_el = c.find("s:v", s_ns)
                if c_type == "inlineStr":
                    t_el = c.find("s:is/s:t", s_ns)
                    val = t_el.text if t_el is not None else ""
                elif c_type == "s":
                    if v_el is None or v_el.text is None:
                        val = ""
                    else:
                        sid = int(v_el.text)
                        val = shared[sid] if 0 <= sid < len(shared) else ""
                else:
                    if v_el is None or v_el.text is None:
                        val = ""
                    else:
                        txt = v_el.text
                        try:
                            if "." in txt:
                                val = float(txt)
                            else:
                                val = int(txt)
                        except ValueError:
                            val = txt
                rmap[c_idx] = val

        if not rows:
            return pd.DataFrame()

        max_row = max(rows)
        max_col = max((max(cols.keys()) for cols in rows.values() if cols), default=0)
        matrix: list[list[object]] = []
        for r in range(max_row + 1):
            cmap = rows.get(r, {})
            matrix.append([cmap.get(c, "") for c in range(max_col + 1)])
        return pd.DataFrame(matrix)


def read_excel_raw_sheet(path: Path) -> pd.DataFrame | None:
    ext = path.suffix.lower()
    if ext == ".xls":
        try:
            return _parse_biff_xls_raw(path)
        except Exception:
            return None
    if ext != ".xlsx":
        return None
    try:
        return pd.read_excel(path, sheet_name=0, header=None)
    except Exception:
        try:
            return _read_strict_xlsx_raw(path)
        except Exception:
            return None


def read_excel_flexible(path: Path) -> pd.DataFrame:
    ext = path.suffix.lower()
    if ext == ".xls":
        return _parse_biff_xls(path)
    if ext == ".xlsx":
        try:
            return pd.read_excel(path, sheet_name=0)
        except Exception:
            return _read_strict_xlsx(path)
    raise ValueError(f"지원하지 않는 파일 형식: {path.name}")


def _postprocess_normalized(out: pd.DataFrame, source_file: str) -> pd.DataFrame:
    if out is None or out.empty:
        out = pd.DataFrame(columns=TARGET_COLUMNS)
    for col in TARGET_COLUMNS:
        if col not in out.columns:
            out[col] = pd.NA
    out = out[TARGET_COLUMNS].copy()
    out["출처파일"] = source_file

    out["사용처"] = out["사용처"].fillna("").astype(str).str.strip()
    out["카테고리"] = out["카테고리"].fillna("").astype(str).str.strip()
    out["세부 카테고리"] = out["세부 카테고리"].fillna("").astype(str).str.strip()
    out["거래구분"] = out["거래구분"].fillna("").astype(str).str.strip()
    out["지출금액"] = pd.to_numeric(out["지출금액"], errors="coerce")
    out["지출날짜"] = pd.to_datetime(out["지출날짜"], errors="coerce")

    # 은행계좌에서 빠져나가는 월별 카드결제 대금은 지출 분석에서 제외 처리.
    settlement_mask = out["사용처"].apply(lambda x: is_monthly_card_settlement(str(x), source_file))
    if settlement_mask.any():
        out.loc[settlement_mask, "카테고리"] = "제외"
        out.loc[settlement_mask, "세부 카테고리"] = "카드결제(월납)"

    out = out[out["사용처"] != ""]
    out = out[out["지출금액"].notna() & (out["지출금액"] > 0)]
    out = out[out["지출날짜"].notna()]
    return out.reset_index(drop=True)


def apply_column_mapping(df: pd.DataFrame, source_file: str, column_map: dict[str, str]) -> pd.DataFrame:
    out = pd.DataFrame(index=df.index)
    for target in TARGET_COLUMNS:
        src = str(column_map.get(target, "")).strip()
        if src and src in df.columns:
            out[target] = df[src]
        else:
            out[target] = pd.NA
    return _postprocess_normalized(out, source_file)


def normalize_columns(df: pd.DataFrame, source_file: str) -> pd.DataFrame:
    alias_map = infer_alias_mapping(list(df.columns))
    return apply_column_mapping(df, source_file, alias_map)


def normalize_columns_with_profile(raw_df: pd.DataFrame, source_file: str, profile: dict) -> pd.DataFrame:
    header_row = int(profile.get("header_row", 0))
    column_map = profile.get("column_map", {})
    shaped = build_dataframe_from_raw(raw_df, header_row)
    return apply_column_mapping(shaped, source_file, column_map if isinstance(column_map, dict) else {})


def find_matching_profile(profiles: list[dict], source_name: str, raw_df: pd.DataFrame | None) -> dict | None:
    if raw_df is None or raw_df.empty:
        return None
    source_norm = _normalize_text(source_name)
    for p in profiles:
        profile = _normalize_profile(p)
        if profile is None:
            continue
        keyword = _normalize_text(profile.get("source_keyword", ""))
        if keyword and keyword not in source_norm:
            continue
        validated = validate_mapping_profile(raw_df, profile)
        if validated is not None:
            return validated
    return None


def _upsert_mapping_profile(profiles: list[dict], profile: dict) -> None:
    norm = _normalize_profile(profile)
    if norm is None:
        return
    key = norm["profile_name"]
    for i, p in enumerate(profiles):
        prev = _normalize_profile(p)
        if prev and prev["profile_name"] == key:
            profiles[i] = norm
            return
    profiles.append(norm)


def load_and_merge(
    files: list[Path],
    logger,
    mapping_profiles: list[dict] | None = None,
    mapping_prompt=None,
    ai_mapping_suggester=None,
) -> pd.DataFrame:
    profiles = mapping_profiles if mapping_profiles is not None else []
    parts: list[pd.DataFrame] = []
    for p in files:
        logger(f"읽는 중: {p.name}")
        raw_sheet = read_excel_raw_sheet(p)
        try:
            raw = read_excel_flexible(p)
        except Exception as e:
            if raw_sheet is None or raw_sheet.empty:
                raise
            guessed_header_row = guess_header_row(raw_sheet)
            logger(
                f"  - 표준 파서 실패({type(e).__name__}: {e}) -> "
                f"헤더 추정 {guessed_header_row + 1}행으로 재시도"
            )
            raw = build_dataframe_from_raw(raw_sheet, guessed_header_row)

        alias_map = infer_alias_mapping(list(raw.columns))
        missing_required = [t for t in REQUIRED_TARGET_COLUMNS if t not in alias_map]
        norm = normalize_columns(raw, p.name)
        profile_used = None

        if missing_required:
            logger(f"  - 기본 컬럼 매핑 누락: {', '.join(missing_required)}")
            if raw_sheet is None or raw_sheet.empty:
                raw_sheet = dataframe_to_raw_like(raw)

            matched = find_matching_profile(profiles, p.name, raw_sheet)
            if matched is not None:
                norm = normalize_columns_with_profile(raw_sheet, p.name, matched)
                profile_used = matched
                logger(f"  - 저장된 매핑 프로필 적용: {matched.get('profile_name', '')}")
            else:
                if ai_mapping_suggester is not None:
                    ai_profile = ai_mapping_suggester.suggest_profile(p, raw_sheet, missing_required)
                    if ai_profile is not None:
                        norm = normalize_columns_with_profile(raw_sheet, p.name, ai_profile)
                        profile_used = ai_profile
                        _upsert_mapping_profile(profiles, ai_profile)
                        logger(f"  - GPT 포맷 매핑 적용: {ai_profile.get('profile_name', '')}")
                    else:
                        logger("  - GPT 포맷 매핑 실패(수동 매핑으로 전환)")

                if profile_used is None and mapping_prompt is not None:
                    manual_profile = mapping_prompt(p, raw_sheet)
                    if manual_profile is None:
                        raise ValueError(f"{p.name}: 컬럼 매핑이 필요합니다.")
                    norm = normalize_columns_with_profile(raw_sheet, p.name, manual_profile)

        parts.append(norm)
        logger(f"  - {len(norm)}건 로드")
        if profile_used is not None and len(norm) == 0:
            logger("  - 경고: 프로필 적용 후에도 유효 지출 건수가 0건입니다.")
        auto_excluded = (
            (norm["카테고리"].fillna("").astype(str).str.strip() == "제외")
            & (norm["세부 카테고리"].fillna("").astype(str).str.strip() == "카드결제(월납)")
        ).sum()
        if auto_excluded:
            logger(f"  - 카드결제 자동 제외 처리: {int(auto_excluded)}건")

    if not parts:
        raise ValueError("읽을 수 있는 파일이 없습니다.")

    merged = pd.concat(parts, ignore_index=True)
    merged = merged.sort_values(["지출날짜", "사용처"], na_position="last").reset_index(drop=True)
    return merged


def load_memory() -> pd.DataFrame:
    if not MEMORY_PATH.exists():
        return pd.DataFrame(columns=["사용처", "카테고리", "세부 카테고리", "count"])
    mem = pd.read_csv(MEMORY_PATH)
    required = {"사용처", "카테고리", "세부 카테고리", "count"}
    if not required.issubset(mem.columns):
        return pd.DataFrame(columns=["사용처", "카테고리", "세부 카테고리", "count"])
    mem["사용처"] = mem["사용처"].fillna("").astype(str).str.strip()
    mem["카테고리"] = mem["카테고리"].fillna("").astype(str).str.strip()
    mem["세부 카테고리"] = mem["세부 카테고리"].fillna("").astype(str).str.strip()
    mem["count"] = pd.to_numeric(mem["count"], errors="coerce").fillna(0)
    mem = mem[mem["사용처"] != ""]
    mem = mem[mem["카테고리"] != ""]
    mem = mem[mem["count"] > 0]
    return mem.reset_index(drop=True)


def update_memory(labeled_df: pd.DataFrame) -> None:
    cur = labeled_df.copy()
    cur["사용처"] = cur["사용처"].fillna("").astype(str).str.strip()
    cur["카테고리"] = cur["카테고리"].fillna("").astype(str).str.strip()
    cur["세부 카테고리"] = cur["세부 카테고리"].fillna("").astype(str).str.strip()
    cur = cur[cur["사용처"] != ""]
    cur = cur[cur["카테고리"] != ""]
    cur = cur[~cur["카테고리"].isin(EXCLUDED_CATEGORIES)]

    grouped = (
        cur.groupby(["사용처", "카테고리", "세부 카테고리"], as_index=False)
        .size()
        .rename(columns={"size": "count"})
    )
    existing = load_memory()
    merged = pd.concat([existing, grouped], ignore_index=True)
    merged = (
        merged.groupby(["사용처", "카테고리", "세부 카테고리"], as_index=False)["count"]
        .sum()
        .sort_values(["사용처", "count"], ascending=[True, False])
    )
    merged.to_csv(MEMORY_PATH, index=False)


def _pick_top(weight_map: dict[tuple[str, str], float]) -> tuple[tuple[str, str], float]:
    items = sorted(weight_map.items(), key=lambda x: x[1], reverse=True)
    top_pair, top_weight = items[0]
    total = sum(v for _, v in items)
    ratio = (top_weight / total) if total > 0 else 0.0
    return top_pair, ratio


def build_knowledge(df: pd.DataFrame, ontology: dict[str, list[str]], logger) -> KnowledgeBase:
    labeled = df.copy()
    labeled["카테고리"] = labeled["카테고리"].fillna("").astype(str).str.strip()
    labeled["세부 카테고리"] = labeled["세부 카테고리"].fillna("").astype(str).str.strip()
    labeled["사용처"] = labeled["사용처"].fillna("").astype(str).str.strip()
    labeled = labeled[labeled["사용처"] != ""]
    labeled = labeled[labeled["카테고리"] != ""]

    allowed = normalize_ontology(ontology)
    mem = load_memory()
    logger(f"기존 메모리 로드: {len(mem)}건")

    merchant_map: dict[str, dict[tuple[str, str], float]] = {}
    keyword_map: dict[str, dict[tuple[str, str], float]] = {}
    subcat_map: dict[str, set[str]] = {}

    for cat, subs in allowed.items():
        subcat_map.setdefault(cat, set()).update(subs if subs else {"미분류"})

    def add_record(merchant: str, category: str, subcategory: str, weight: float) -> None:
        if not merchant or not category:
            return
        if category not in allowed:
            return
        if category.lower() == "nan":
            return
        allowed_subs = allowed.get(category, [])
        if not subcategory or subcategory.lower() == "nan":
            subcategory = "미분류"
        if allowed_subs and subcategory not in allowed_subs:
            subcategory = "미분류" if "미분류" in allowed_subs else allowed_subs[0]
        key = (category, subcategory)
        merchant_map.setdefault(merchant, {})
        merchant_map[merchant][key] = merchant_map[merchant].get(key, 0.0) + weight
        for tok in tokenize(merchant):
            keyword_map.setdefault(tok, {})
            keyword_map[tok][key] = keyword_map[tok].get(key, 0.0) + weight
        subcat_map.setdefault(category, set()).add(subcategory if subcategory else "미분류")

    for row in labeled[["사용처", "카테고리", "세부 카테고리"]].itertuples(index=False, name=None):
        add_record(str(row[0]).strip(), str(row[1]).strip(), str(row[2]).strip(), 1.0)

    for row in mem[["사용처", "카테고리", "세부 카테고리", "count"]].itertuples(index=False, name=None):
        add_record(str(row[0]).strip(), str(row[1]).strip(), str(row[2]).strip(), float(row[3]))

    subcat_by_category = {k: sorted(v) for k, v in subcat_map.items()}
    return KnowledgeBase(merchant_map=merchant_map, keyword_map=keyword_map, subcat_by_category=subcat_by_category)


def suggest_from_rules(merchant: str, kb: KnowledgeBase) -> TagSuggestion | None:
    merchant = (merchant or "").strip()
    if not merchant:
        return None

    if merchant in kb.merchant_map:
        (cat, sub), ratio = _pick_top(kb.merchant_map[merchant])
        conf = min(0.99, 0.85 + 0.14 * ratio)
        return TagSuggestion(cat, sub, conf, "rule_exact", "동일 사용처 이력 기반")

    combined: dict[tuple[str, str], float] = {}
    for tok in tokenize(merchant):
        for key, val in kb.keyword_map.get(tok, {}).items():
            combined[key] = combined.get(key, 0.0) + val

    if combined:
        (cat, sub), ratio = _pick_top(combined)
        conf = min(0.79, 0.45 + 0.34 * ratio)
        return TagSuggestion(cat, sub, conf, "rule_keyword", "사용처 키워드 유사도 기반")
    return None


class OpenAIColumnMapper:
    def __init__(self, api_key: str, model: str = "gpt-4o-mini", timeout: int = 35):
        self.api_key = api_key.strip()
        self.model = model.strip() or "gpt-4o-mini"
        self.timeout = timeout

    @staticmethod
    def _build_preview(raw_df: pd.DataFrame, max_rows: int = 30, max_cols: int = 24) -> list[list[str]]:
        if raw_df is None or raw_df.empty:
            return []
        clipped = raw_df.iloc[:max_rows, :max_cols].fillna("")
        preview: list[list[str]] = []
        for _, row in clipped.iterrows():
            preview.append([str(v).strip() for v in row.tolist()])
        return preview

    def suggest_profile(self, source_path: Path, raw_df: pd.DataFrame, missing_required: list[str] | None = None) -> dict | None:
        if not self.api_key or raw_df is None or raw_df.empty:
            return None

        preview = self._build_preview(raw_df)
        if not preview:
            return None

        system_prompt = (
            "너는 거래내역 엑셀 포맷 분석기다. 반드시 JSON만 출력한다. "
            "반환 필드: header_row(number, 0-based), column_map(object). "
            "column_map은 표준 타겟 컬럼명을 key로, 원본 헤더 문자열을 value로 넣어라."
        )
        user_payload = {
            "source_file": source_path.name,
            "target_columns": TARGET_COLUMNS,
            "required_target_columns": REQUIRED_TARGET_COLUMNS,
            "currently_missing_required": missing_required or [],
            "alias_examples": get_column_aliases(),
            "raw_sheet_preview": preview,
            "rules": [
                "header_row는 0부터 시작한다.",
                "column_map value는 반드시 header_row의 셀에 실제로 존재하는 문자열을 사용한다.",
                "확신 없는 매핑은 생략한다.",
            ],
        }
        user_prompt = json.dumps(user_payload, ensure_ascii=False)

        payload = {
            "model": self.model,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            "temperature": 0.0,
            "response_format": {"type": "json_object"},
        }

        req = urllib.request.Request(
            "https://api.openai.com/v1/chat/completions",
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {self.api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )

        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                body = json.loads(resp.read().decode("utf-8"))
            content = body["choices"][0]["message"]["content"]
            data = json.loads(content)
        except Exception:
            return None

        if not isinstance(data, dict):
            return None
        try:
            header_row = int(data.get("header_row", 0))
        except Exception:
            header_row = 0

        raw_map = data.get("column_map", {})
        if not isinstance(raw_map, dict):
            raw_map = {}

        normalized_key_map = {_normalize_text(k).replace(" ", ""): v for k, v in raw_map.items()}
        column_map: dict[str, str] = {}
        for target in TARGET_COLUMNS:
            src = raw_map.get(target, "")
            if not src:
                src = normalized_key_map.get(_normalize_text(target).replace(" ", ""), "")
            src_text = str(src).strip()
            if src_text:
                column_map[target] = src_text

        profile_name = f"gpt_auto_{safe_filename(source_path.stem) or 'mapping'}"
        draft_profile = {
            "profile_name": profile_name,
            "source_keyword": source_path.stem,
            "header_row": max(0, header_row),
            "column_map": column_map,
        }
        validated = validate_mapping_profile(raw_df, draft_profile)
        if validated is None and header_row > 0:
            draft_profile["header_row"] = header_row - 1
            validated = validate_mapping_profile(raw_df, draft_profile)
        return validated


class OpenAITagger:
    def __init__(self, api_key: str, model: str = "gpt-4o-mini", timeout: int = 25):
        self.api_key = api_key.strip()
        self.model = model.strip() or "gpt-4o-mini"
        self.timeout = timeout

    def classify(
        self,
        merchant: str,
        amount: float,
        date_text: str,
        taxonomy: dict[str, list[str]],
    ) -> TagSuggestion | None:
        if not self.api_key:
            return None

        tax_lines = []
        for cat, subcats in taxonomy.items():
            sample = ", ".join(subcats[:10]) if subcats else "미분류"
            tax_lines.append(f"- {cat}: {sample}")
        taxonomy_text = "\n".join(tax_lines) if tax_lines else "- (분류 후보 없음)"

        system_prompt = (
            "너는 가계부 거래 내역 분류기다. 반드시 JSON만 출력한다. "
            '필드: category(string), subcategory(string), confidence(number 0~1), reason(string). '
            "category/subcategory는 한국어로 작성한다."
        )
        user_prompt = (
            f"거래일: {date_text}\n"
            f"사용처: {merchant}\n"
            f"금액: {amount:,.0f}\n\n"
            f"가능한 카테고리/세부카테고리 참고:\n{taxonomy_text}\n\n"
            "가장 적합한 카테고리와 세부카테고리를 추정해줘."
        )

        payload = {
            "model": self.model,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            "temperature": 0.1,
            "response_format": {"type": "json_object"},
        }

        req = urllib.request.Request(
            "https://api.openai.com/v1/chat/completions",
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {self.api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )

        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                body = json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            return TagSuggestion("", "", 0.0, "ai_error", f"HTTPError: {e.code}")
        except Exception as e:
            return TagSuggestion("", "", 0.0, "ai_error", f"{type(e).__name__}: {e}")

        try:
            content = body["choices"][0]["message"]["content"]
            data = json.loads(content)
        except Exception:
            return TagSuggestion("", "", 0.0, "ai_error", "응답 파싱 실패")

        category = str(data.get("category", "")).strip()
        subcategory = str(data.get("subcategory", "")).strip() or "미분류"
        reason = str(data.get("reason", "")).strip()
        try:
            confidence = float(data.get("confidence", 0.0))
        except Exception:
            confidence = 0.0
        confidence = max(0.0, min(1.0, confidence))

        return TagSuggestion(category, subcategory, confidence, "ai", reason)


class ColumnMappingDialog(tk.Toplevel):
    def __init__(self, parent: tk.Tk, source_path: Path, raw_df: pd.DataFrame):
        super().__init__(parent)
        self.source_path = source_path
        self.raw_df = raw_df.copy()
        self.result: dict | None = None
        self.title(f"컬럼 매핑 - {source_path.name}")
        self.geometry("980x620")
        self.transient(parent)
        self.grab_set()

        self.none_label = "(없음)"
        self.header_row_var = tk.StringVar()
        self.save_profile_var = tk.BooleanVar(value=True)
        self.profile_name_var = tk.StringVar(value=f"{source_path.stem}_매핑")
        self.source_keyword_var = tk.StringVar(value=source_path.stem)
        self.mapping_vars = {t: tk.StringVar(value=self.none_label) for t in TARGET_COLUMNS}
        self.mapping_combos: dict[str, ttk.Combobox] = {}

        self._build_ui()
        guess = self._guess_header_row()
        self.header_row_var.set(str(guess + 1))
        self._refresh_after_header_change()
        self.protocol("WM_DELETE_WINDOW", self._cancel)

    def _build_ui(self) -> None:
        root = ttk.Frame(self, padding=10)
        root.pack(fill="both", expand=True)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(2, weight=1)

        info = (
            "이 파일 포맷의 헤더 행을 고르고, 표준 컬럼에 대응되는 원본 컬럼을 지정하세요.\n"
            "필수: 지출날짜, 사용처, 지출금액"
        )
        ttk.Label(root, text=info, justify="left").grid(row=0, column=0, sticky="w")

        top = ttk.Frame(root)
        top.grid(row=1, column=0, sticky="ew", pady=(8, 8))
        top.columnconfigure(2, weight=1)

        ttk.Label(top, text="헤더 행(1부터)").grid(row=0, column=0, sticky="w")
        max_rows = max(1, min(40, len(self.raw_df)))
        self.header_combo = ttk.Combobox(top, textvariable=self.header_row_var, values=[str(i) for i in range(1, max_rows + 1)], width=8)
        self.header_combo.grid(row=0, column=1, sticky="w", padx=(6, 12))
        self.header_combo.bind("<<ComboboxSelected>>", lambda _e: self._refresh_after_header_change())
        self.header_combo.bind("<FocusOut>", lambda _e: self._refresh_after_header_change())

        map_frame = ttk.LabelFrame(root, text="컬럼 매핑")
        map_frame.grid(row=2, column=0, sticky="nsew")
        map_frame.columnconfigure(1, weight=1)

        for i, target in enumerate(TARGET_COLUMNS):
            mark = " *" if target in REQUIRED_TARGET_COLUMNS else ""
            ttk.Label(map_frame, text=f"{target}{mark}").grid(row=i, column=0, sticky="w", padx=8, pady=4)
            combo = ttk.Combobox(map_frame, textvariable=self.mapping_vars[target], values=[self.none_label], width=60)
            combo.grid(row=i, column=1, sticky="ew", padx=(0, 8), pady=4)
            self.mapping_combos[target] = combo

        preview_frame = ttk.LabelFrame(root, text="미리보기")
        preview_frame.grid(row=3, column=0, sticky="nsew", pady=(8, 0))
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(0, weight=1)
        self.preview_text = tk.Text(preview_frame, height=12, wrap="none")
        self.preview_text.grid(row=0, column=0, sticky="nsew", padx=8, pady=8)

        profile_frame = ttk.LabelFrame(root, text="프로필 저장")
        profile_frame.grid(row=4, column=0, sticky="ew", pady=(8, 0))
        profile_frame.columnconfigure(1, weight=1)
        ttk.Checkbutton(profile_frame, text="이 매핑을 프로필로 저장", variable=self.save_profile_var).grid(
            row=0, column=0, columnspan=2, sticky="w", padx=8, pady=(8, 4)
        )
        ttk.Label(profile_frame, text="프로필 이름").grid(row=1, column=0, sticky="w", padx=8, pady=2)
        ttk.Entry(profile_frame, textvariable=self.profile_name_var).grid(row=1, column=1, sticky="ew", padx=(0, 8), pady=2)
        ttk.Label(profile_frame, text="파일명 키워드").grid(row=2, column=0, sticky="w", padx=8, pady=(2, 8))
        ttk.Entry(profile_frame, textvariable=self.source_keyword_var).grid(row=2, column=1, sticky="ew", padx=(0, 8), pady=(2, 8))

        btns = ttk.Frame(root)
        btns.grid(row=5, column=0, sticky="e", pady=(10, 0))
        ttk.Button(btns, text="적용", command=self._apply).grid(row=0, column=0, padx=4)
        ttk.Button(btns, text="취소", command=self._cancel).grid(row=0, column=1, padx=4)

    def _guess_header_row(self) -> int:
        return guess_header_row(self.raw_df)

    def _current_header_row(self) -> int:
        try:
            idx = int(self.header_row_var.get().strip()) - 1
        except Exception:
            idx = 0
        return max(0, min(idx, max(0, len(self.raw_df) - 1)))

    def _refresh_after_header_change(self) -> None:
        header_row = self._current_header_row()
        self.header_row_var.set(str(header_row + 1))
        shaped = build_dataframe_from_raw(self.raw_df, header_row)
        cols = list(shaped.columns)
        options = [self.none_label] + cols

        inferred = infer_alias_mapping(cols)
        for target, combo in self.mapping_combos.items():
            combo["values"] = options
            cur = self.mapping_vars[target].get().strip()
            if cur not in options:
                guess = inferred.get(target, self.none_label)
                self.mapping_vars[target].set(guess if guess in options else self.none_label)

        self.preview_text.delete("1.0", "end")
        if shaped.empty:
            self.preview_text.insert("end", "(미리보기 데이터가 없습니다)")
            return
        preview = shaped.iloc[:8, : min(10, shaped.shape[1])]
        self.preview_text.insert("end", preview.to_string(index=False))

    def _apply(self) -> None:
        for target in REQUIRED_TARGET_COLUMNS:
            if self.mapping_vars[target].get().strip() in ("", self.none_label):
                messagebox.showwarning("확인", f"{target} 컬럼을 지정해 주세요.", parent=self)
                return

        if self.save_profile_var.get():
            if not self.profile_name_var.get().strip():
                messagebox.showwarning("확인", "프로필 이름을 입력해 주세요.", parent=self)
                return

        column_map: dict[str, str] = {}
        for target in TARGET_COLUMNS:
            src = self.mapping_vars[target].get().strip()
            if src and src != self.none_label:
                column_map[target] = src

        self.result = {
            "profile_name": self.profile_name_var.get().strip() or f"{self.source_path.stem}_매핑",
            "source_keyword": self.source_keyword_var.get().strip(),
            "header_row": self._current_header_row(),
            "column_map": column_map,
            "save_profile": bool(self.save_profile_var.get()),
        }
        self.destroy()

    def _cancel(self) -> None:
        self.result = None
        self.destroy()


class ReviewDialog(tk.Toplevel):
    def __init__(
        self,
        parent: tk.Tk,
        row: pd.Series,
        suggestion: TagSuggestion | None,
        ontology: dict[str, list[str]],
    ):
        super().__init__(parent)
        self.result: tuple[str, str] | None | str = None
        self.ontology = ontology
        self.ontology_changed = False
        self.title("불확실 항목 검토")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()

        frm = ttk.Frame(self, padding=12)
        frm.grid(row=0, column=0, sticky="nsew")

        info = (
            f"날짜: {row['지출날짜'].strftime('%Y-%m-%d') if pd.notna(row['지출날짜']) else ''}\n"
            f"사용처: {row['사용처']}\n"
            f"금액: {row['지출금액']:,.0f}원\n"
            f"현재 카테고리: {row['카테고리'] or '(없음)'} / {row['세부 카테고리'] or '(없음)'}"
        )
        ttk.Label(frm, text=info, justify="left").grid(row=0, column=0, columnspan=2, sticky="w")

        if suggestion:
            sug = (
                f"추천: {suggestion.category} / {suggestion.subcategory} "
                f"(신뢰도 {suggestion.confidence:.2f}, {suggestion.source})"
            )
            ttk.Label(frm, text=sug, foreground="#1f4f8a").grid(row=1, column=0, columnspan=2, sticky="w", pady=(8, 0))

        ttk.Label(frm, text="카테고리").grid(row=2, column=0, sticky="w", pady=(12, 2))
        init_cat = suggestion.category if suggestion else str(row.get("카테고리", "")).strip()
        init_sub = suggestion.subcategory if suggestion else str(row.get("세부 카테고리", "")).strip()
        self.cat_var = tk.StringVar(value=init_cat)
        self.sub_var = tk.StringVar(value=init_sub)

        self.cat_combo = ttk.Combobox(frm, textvariable=self.cat_var, values=[], width=34)
        self.cat_combo.grid(row=3, column=0, sticky="ew")
        self.cat_combo.bind("<<ComboboxSelected>>", self._on_cat_change)
        cat_btn = ttk.Frame(frm)
        cat_btn.grid(row=3, column=1, sticky="e")
        ttk.Button(cat_btn, text="카테고리 추가", command=self._add_category).grid(row=0, column=0, padx=(6, 0))
        ttk.Button(cat_btn, text="카테고리 삭제", command=self._delete_category).grid(row=0, column=1, padx=(4, 0))

        ttk.Label(frm, text="세부 카테고리").grid(row=4, column=0, sticky="w", pady=(8, 2))
        self.sub_combo = ttk.Combobox(frm, textvariable=self.sub_var, values=[], width=34)
        self.sub_combo.grid(row=5, column=0, sticky="ew")
        sub_btn = ttk.Frame(frm)
        sub_btn.grid(row=5, column=1, sticky="e")
        ttk.Button(sub_btn, text="세부 추가", command=self._add_subcategory).grid(row=0, column=0, padx=(6, 0))
        ttk.Button(sub_btn, text="세부 삭제", command=self._delete_subcategory).grid(row=0, column=1, padx=(4, 0))

        btn_frame = ttk.Frame(frm)
        btn_frame.grid(row=6, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(btn_frame, text="적용", command=self._apply).grid(row=0, column=0, padx=4)
        ttk.Button(btn_frame, text="제외 처리", command=self._exclude).grid(row=0, column=1, padx=4)
        ttk.Button(btn_frame, text="건너뛰기", command=self._skip).grid(row=0, column=2, padx=4)
        ttk.Button(btn_frame, text="검토 중단", command=self._cancel_all).grid(row=0, column=3, padx=4)

        self._refresh_category_values()
        self._on_cat_change()
        self.protocol("WM_DELETE_WINDOW", self._skip)

    def _category_values(self) -> list[str]:
        vals = set(self.ontology.keys()) | set(EXCLUDED_CATEGORIES)
        current = self.cat_var.get().strip()
        if current:
            vals.add(current)
        return sorted(vals)

    def _refresh_category_values(self, keep: str | None = None) -> None:
        values = self._category_values()
        self.cat_combo["values"] = values
        target = (keep or self.cat_var.get().strip())
        if not target and values:
            target = values[0]
        if target:
            self.cat_var.set(target)

    def _on_cat_change(self, _event=None) -> None:
        cat = self.cat_var.get().strip()
        options = list(self.ontology.get(cat, []))
        current_sub = self.sub_var.get().strip()
        if current_sub and current_sub not in options:
            options.append(current_sub)
        if not options:
            options = ["미분류"]
        self.sub_combo["values"] = options
        if self.sub_var.get().strip() not in options:
            self.sub_var.set(options[0] if options else "미분류")

    def _add_category(self) -> None:
        name = simpledialog.askstring("카테고리 추가", "새 카테고리 이름", parent=self)
        if not name:
            return
        category = name.strip()
        if not category:
            return
        if category in self.ontology:
            messagebox.showwarning("확인", "이미 존재하는 카테고리입니다.", parent=self)
            return
        self.ontology[category] = ["미분류"]
        self.ontology_changed = True
        self._refresh_category_values(keep=category)
        self._on_cat_change()

    def _delete_category(self) -> None:
        category = self.cat_var.get().strip()
        if not category:
            return
        if category in EXCLUDED_CATEGORIES:
            messagebox.showwarning("확인", f'"{category}" 카테고리는 삭제할 수 없습니다.', parent=self)
            return
        if category not in self.ontology:
            messagebox.showwarning("확인", "온톨로지에 없는 카테고리입니다.", parent=self)
            return
        ok = messagebox.askyesno("확인", f'"{category}" 카테고리를 삭제할까요?', parent=self)
        if not ok:
            return
        self.ontology.pop(category, None)
        self.ontology_changed = True
        self._refresh_category_values()
        self._on_cat_change()

    def _add_subcategory(self) -> None:
        category = self.cat_var.get().strip()
        if not category:
            messagebox.showwarning("확인", "먼저 카테고리를 선택해 주세요.", parent=self)
            return
        name = simpledialog.askstring("세부 카테고리 추가", "새 세부 카테고리 이름", parent=self)
        if not name:
            return
        sub = name.strip()
        if not sub:
            return
        self.ontology.setdefault(category, [])
        if sub in self.ontology[category]:
            messagebox.showwarning("확인", "이미 존재하는 세부 카테고리입니다.", parent=self)
            return
        self.ontology[category].append(sub)
        self.ontology_changed = True
        self.sub_var.set(sub)
        self._on_cat_change()

    def _delete_subcategory(self) -> None:
        category = self.cat_var.get().strip()
        sub = self.sub_var.get().strip()
        if not category or not sub:
            return
        if category not in self.ontology:
            messagebox.showwarning("확인", "온톨로지에 없는 카테고리입니다.", parent=self)
            return
        if sub not in self.ontology[category]:
            messagebox.showwarning("확인", "온톨로지에 없는 세부 카테고리입니다.", parent=self)
            return
        ok = messagebox.askyesno("확인", f'"{sub}" 세부 카테고리를 삭제할까요?', parent=self)
        if not ok:
            return
        subs = [x for x in self.ontology[category] if x != sub]
        if not subs:
            subs = ["미분류"]
        self.ontology[category] = subs
        self.ontology_changed = True
        self.sub_var.set(subs[0])
        self._on_cat_change()

    def _apply(self) -> None:
        cat = self.cat_var.get().strip()
        sub = self.sub_var.get().strip() or "미분류"
        if not cat:
            messagebox.showwarning("확인", "카테고리를 입력해 주세요.", parent=self)
            return
        self.result = (cat, sub)
        self.destroy()

    def _exclude(self) -> None:
        self.result = ("제외", "수동 제외")
        self.destroy()

    def _skip(self) -> None:
        self.result = None
        self.destroy()

    def _cancel_all(self) -> None:
        self.result = "__CANCEL_ALL__"
        self.destroy()


class OntologyEditorDialog(tk.Toplevel):
    def __init__(self, parent: tk.Tk, ontology: dict[str, list[str]]):
        super().__init__(parent)
        self.title("온톨로지 편집")
        self.geometry("760x460")
        self.resizable(True, True)
        self.transient(parent)
        self.grab_set()
        self.result: dict[str, list[str]] | None = None
        self.ontology: dict[str, list[str]] = {
            c: list(s) for c, s in normalize_ontology(ontology).items()
        }

        self._build_ui()
        self._refresh_categories()
        self.protocol("WM_DELETE_WINDOW", self._cancel)

    def _build_ui(self) -> None:
        root = ttk.Frame(self, padding=10)
        root.pack(fill="both", expand=True)
        root.columnconfigure(0, weight=1)
        root.columnconfigure(1, weight=1)
        root.rowconfigure(1, weight=1)

        ttk.Label(root, text="카테고리").grid(row=0, column=0, sticky="w")
        ttk.Label(root, text="세부 카테고리").grid(row=0, column=1, sticky="w")

        self.cat_list = tk.Listbox(root, exportselection=False)
        self.cat_list.grid(row=1, column=0, sticky="nsew", padx=(0, 8))
        self.cat_list.bind("<<ListboxSelect>>", self._on_category_select)

        self.sub_list = tk.Listbox(root, exportselection=False)
        self.sub_list.grid(row=1, column=1, sticky="nsew")

        cat_btn = ttk.Frame(root)
        cat_btn.grid(row=2, column=0, sticky="ew", pady=(8, 0), padx=(0, 8))
        ttk.Button(cat_btn, text="카테고리 추가", command=self._add_category).pack(side="left")
        ttk.Button(cat_btn, text="카테고리 삭제", command=self._delete_category).pack(side="left", padx=(6, 0))

        sub_btn = ttk.Frame(root)
        sub_btn.grid(row=2, column=1, sticky="ew", pady=(8, 0))
        ttk.Button(sub_btn, text="세부 추가", command=self._add_subcategory).pack(side="left")
        ttk.Button(sub_btn, text="세부 삭제", command=self._delete_subcategory).pack(side="left", padx=(6, 0))

        bottom = ttk.Frame(root)
        bottom.grid(row=3, column=0, columnspan=2, sticky="e", pady=(10, 0))
        ttk.Button(bottom, text="기본값 복원", command=self._restore_default).pack(side="left", padx=4)
        ttk.Button(bottom, text="저장", command=self._save).pack(side="left", padx=4)
        ttk.Button(bottom, text="취소", command=self._cancel).pack(side="left", padx=4)

    def _selected_category(self) -> str | None:
        sel = self.cat_list.curselection()
        if not sel:
            return None
        return str(self.cat_list.get(sel[0]))

    def _selected_subcategory(self) -> str | None:
        sel = self.sub_list.curselection()
        if not sel:
            return None
        return str(self.sub_list.get(sel[0]))

    def _refresh_categories(self, keep: str | None = None) -> None:
        categories = sorted(self.ontology.keys())
        self.cat_list.delete(0, "end")
        for cat in categories:
            self.cat_list.insert("end", cat)

        if not categories:
            self.sub_list.delete(0, "end")
            return

        target = keep if keep in categories else categories[0]
        idx = categories.index(target)
        self.cat_list.selection_set(idx)
        self.cat_list.activate(idx)
        self._refresh_subcategories(target)

    def _refresh_subcategories(self, category: str | None) -> None:
        self.sub_list.delete(0, "end")
        if not category:
            return
        subs = self.ontology.get(category, [])
        for sub in subs:
            self.sub_list.insert("end", sub)
        if subs:
            self.sub_list.selection_set(0)
            self.sub_list.activate(0)

    def _on_category_select(self, _event=None) -> None:
        self._refresh_subcategories(self._selected_category())

    def _add_category(self) -> None:
        name = simpledialog.askstring("카테고리 추가", "새 카테고리 이름", parent=self)
        if not name:
            return
        category = name.strip()
        if not category:
            return
        if category in self.ontology:
            messagebox.showwarning("확인", "이미 존재하는 카테고리입니다.", parent=self)
            return
        self.ontology[category] = ["미분류"]
        self._refresh_categories(keep=category)

    def _delete_category(self) -> None:
        category = self._selected_category()
        if not category:
            return
        ok = messagebox.askyesno("확인", f'"{category}" 카테고리를 삭제할까요?', parent=self)
        if not ok:
            return
        self.ontology.pop(category, None)
        if not self.ontology:
            self.ontology = get_default_ontology()
        self._refresh_categories()

    def _add_subcategory(self) -> None:
        category = self._selected_category()
        if not category:
            messagebox.showwarning("확인", "먼저 카테고리를 선택해 주세요.", parent=self)
            return
        name = simpledialog.askstring("세부 카테고리 추가", "새 세부 카테고리 이름", parent=self)
        if not name:
            return
        sub = name.strip()
        if not sub:
            return
        if sub in self.ontology.get(category, []):
            messagebox.showwarning("확인", "이미 존재하는 세부 카테고리입니다.", parent=self)
            return
        self.ontology.setdefault(category, []).append(sub)
        self._refresh_subcategories(category)

    def _delete_subcategory(self) -> None:
        category = self._selected_category()
        sub = self._selected_subcategory()
        if not category or not sub:
            return
        ok = messagebox.askyesno("확인", f'"{sub}" 세부 카테고리를 삭제할까요?', parent=self)
        if not ok:
            return
        subs = [x for x in self.ontology.get(category, []) if x != sub]
        if not subs:
            subs = ["미분류"]
        self.ontology[category] = subs
        self._refresh_subcategories(category)

    def _restore_default(self) -> None:
        ok = messagebox.askyesno("확인", "온톨로지를 기본값으로 되돌릴까요?", parent=self)
        if not ok:
            return
        self.ontology = get_default_ontology()
        self._refresh_categories()

    def _save(self) -> None:
        self.result = normalize_ontology(self.ontology)
        self.destroy()

    def _cancel(self) -> None:
        self.result = None
        self.destroy()


def auto_tag(
    df: pd.DataFrame,
    kb: KnowledgeBase,
    ai_tagger: OpenAITagger | None,
    ai_threshold: float,
    logger,
) -> tuple[pd.DataFrame, list[int]]:
    out = df.copy()
    out["추천카테고리"] = ""
    out["추천세부카테고리"] = ""
    out["자동태깅근거"] = ""
    out["신뢰도"] = pd.NA

    taxonomy = kb.subcat_by_category
    allowed_categories = set(taxonomy.keys())
    uncertain: list[int] = []

    def queue_manual_required(idx: int, merchant: str, tx_type: str) -> bool:
        cur_cat = str(out.at[idx, "카테고리"]).strip()
        cur_sub = str(out.at[idx, "세부 카테고리"]).strip()
        if requires_manual_classification(merchant, cur_cat, cur_sub, tx_type):
            out.at[idx, "추천카테고리"] = cur_cat
            out.at[idx, "추천세부카테고리"] = cur_sub or "미분류"
            out.at[idx, "자동태깅근거"] = "review:ambiguous_manual_required"
            prev_conf = pd.to_numeric(pd.Series([out.at[idx, "신뢰도"]]), errors="coerce").iloc[0]
            if pd.isna(prev_conf):
                prev_conf = 0.0
            out.at[idx, "신뢰도"] = round(float(max(0.01, prev_conf)), 3)
            if idx not in uncertain:
                uncertain.append(idx)
            return True
        return False

    for idx, row in out.iterrows():
        cat = str(row["카테고리"]).strip()
        sub = str(row["세부 카테고리"]).strip()
        merchant = str(row["사용처"]).strip()
        tx_type = str(row.get("거래구분", "")).strip()

        if cat in EXCLUDED_CATEGORIES:
            if not sub:
                out.at[idx, "세부 카테고리"] = "미분류"
            out.at[idx, "자동태깅근거"] = "preset_excluded"
            out.at[idx, "신뢰도"] = 1.0
            continue

        if cat and cat.lower() != "nan" and cat not in allowed_categories:
            cat = ""
            sub = ""
            out.at[idx, "카테고리"] = ""
            out.at[idx, "세부 카테고리"] = ""

        if cat and cat.lower() != "nan":
            allowed_subs = taxonomy.get(cat, [])
            if sub and allowed_subs and sub not in allowed_subs:
                sub = ""
                out.at[idx, "세부 카테고리"] = ""
            if not sub:
                rule = suggest_from_rules(merchant, kb)
                if rule and rule.category == cat:
                    out.at[idx, "세부 카테고리"] = rule.subcategory or "미분류"
                    out.at[idx, "자동태깅근거"] = "기존카테고리 + 규칙보완"
                    out.at[idx, "신뢰도"] = round(rule.confidence, 3)
                else:
                    out.at[idx, "세부 카테고리"] = "미분류"
            queue_manual_required(idx, merchant, tx_type)
            continue

        rule = suggest_from_rules(merchant, kb)
        best: TagSuggestion | None = rule
        if rule:
            out.at[idx, "추천카테고리"] = rule.category
            out.at[idx, "추천세부카테고리"] = rule.subcategory

        if rule and rule.confidence >= 0.90:
            out.at[idx, "카테고리"] = rule.category
            out.at[idx, "세부 카테고리"] = rule.subcategory or "미분류"
            out.at[idx, "자동태깅근거"] = rule.source
            out.at[idx, "신뢰도"] = round(rule.confidence, 3)
            queue_manual_required(idx, merchant, tx_type)
            continue

        if ai_tagger is not None:
            date_text = row["지출날짜"].strftime("%Y-%m-%d") if pd.notna(row["지출날짜"]) else ""
            ai = ai_tagger.classify(
                merchant=merchant,
                amount=float(row["지출금액"]),
                date_text=date_text,
                taxonomy=taxonomy,
            )
            if ai and ai.category:
                if (best is None) or (ai.confidence > best.confidence):
                    best = ai
                out.at[idx, "추천카테고리"] = ai.category
                out.at[idx, "추천세부카테고리"] = ai.subcategory
                if ai.confidence >= ai_threshold:
                    out.at[idx, "카테고리"] = ai.category
                    out.at[idx, "세부 카테고리"] = ai.subcategory or "미분류"
                    out.at[idx, "자동태깅근거"] = "ai"
                    out.at[idx, "신뢰도"] = round(ai.confidence, 3)
                    queue_manual_required(idx, merchant, tx_type)
                    continue

        uncertain.append(idx)
        if best:
            out.at[idx, "추천카테고리"] = best.category
            out.at[idx, "추천세부카테고리"] = best.subcategory
            out.at[idx, "자동태깅근거"] = f"review:{best.source}"
            out.at[idx, "신뢰도"] = round(best.confidence, 3)
        else:
            out.at[idx, "자동태깅근거"] = "review:unknown"
            out.at[idx, "신뢰도"] = 0.0

    logger(f"자동태깅 완료: 전체 {len(out)}건, 검토 필요 {len(uncertain)}건")
    return out, uncertain


def spending_only(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["카테고리"] = out["카테고리"].fillna("").astype(str).str.strip()
    out["세부 카테고리"] = out["세부 카테고리"].fillna("").astype(str).str.strip()
    out["세부 카테고리"] = out["세부 카테고리"].replace("", "미분류")
    out = out[out["카테고리"] != ""]
    out = out[out["카테고리"].str.lower() != "nan"]
    out = out[~out["카테고리"].isin(EXCLUDED_CATEGORIES)]
    out = out[out["지출금액"].notna() & (out["지출금액"] > 0)]
    out = out[out["지출날짜"].notna()]
    return out.reset_index(drop=True)


def plot_pie_with_legend(series: pd.Series, title: str, out_path: Path, legend_title: str) -> None:
    if series.empty:
        return
    values = series.values.astype(float)
    labels = series.index.astype(str).tolist()
    total = float(series.sum())
    legend_labels = [
        f"{label}: {(value / total) * 100:.1f}% ({value:,.0f}원)" for label, value in zip(labels, values)
    ]

    fig_height = max(8.0, 4.0 + 0.35 * len(labels))
    fig, ax = plt.subplots(figsize=(15, fig_height))
    wedges, _ = ax.pie(
        values,
        labels=None,
        startangle=90,
        counterclock=False,
        wedgeprops={"linewidth": 1, "edgecolor": "white"},
    )
    ax.legend(
        wedges,
        legend_labels,
        title=legend_title,
        loc="center left",
        bbox_to_anchor=(1.02, 0.5),
        frameon=False,
        fontsize=9,
        title_fontsize=11,
    )
    ax.set_title(f"{title}\n총액: {total:,.0f}원", fontsize=14)
    ax.axis("equal")
    fig.subplots_adjust(right=0.62)
    fig.tight_layout()
    fig.savefig(out_path, dpi=200)
    plt.close(fig)


def export_charts(df_spending: pd.DataFrame, charts_dir: Path, logger) -> list[Path]:
    charts_dir.mkdir(parents=True, exist_ok=True)
    for p in charts_dir.glob("*.png"):
        p.unlink()

    created: list[Path] = []
    cat_sum = df_spending.groupby("카테고리")["지출금액"].sum().sort_values(ascending=False)
    if cat_sum.empty:
        return created

    category_path = charts_dir / "00_카테고리_파이차트.png"
    plot_pie_with_legend(cat_sum, "카테고리별 지출 비중", category_path, "카테고리")
    created.append(category_path)

    for rank, category in enumerate(cat_sum.index, start=1):
        sub_sum = (
            df_spending[df_spending["카테고리"] == category]
            .groupby("세부 카테고리")["지출금액"]
            .sum()
            .sort_values(ascending=False)
        )
        p = charts_dir / f"{rank:02d}_세부카테고리_파이차트_{safe_filename(category)}.png"
        plot_pie_with_legend(sub_sum, f"{category} - 세부 카테고리별 지출 비중", p, "세부 카테고리")
        created.append(p)

    all_sub = df_spending.groupby("세부 카테고리")["지출금액"].sum().sort_values(ascending=False)
    overall_rank = len(cat_sum) + 1
    sub_overall_path = charts_dir / f"{overall_rank:02d}_세부카테고리_전체_파이차트.png"
    plot_pie_with_legend(all_sub, "세부 카테고리 전체 지출 비중", sub_overall_path, "세부 카테고리")
    created.append(sub_overall_path)

    monthly = df_spending.copy()
    monthly["연월"] = monthly["지출날짜"].dt.to_period("M").astype(str)
    for ym in sorted(monthly["연월"].unique()):
        sub = monthly[monthly["연월"] == ym]
        month_cat = sub.groupby("카테고리")["지출금액"].sum().sort_values(ascending=False)
        p = charts_dir / f"월별카테고리_파이차트_{ym}.png"
        plot_pie_with_legend(month_cat, f"{ym} 카테고리별 지출 비중", p, "카테고리")
        created.append(p)

    logger(f"차트 생성 완료: {len(created)}개")
    return created


def export_excels(df_all: pd.DataFrame, df_spending: pd.DataFrame, out_dir: Path, logger) -> tuple[Path, Path]:
    tagged_path = out_dir / "태깅결과_통합지출내역.xlsx"
    monthly_path = out_dir / "월별_총지출_요약.xlsx"

    out_all = df_all.copy()
    out_all["지출날짜"] = out_all["지출날짜"].dt.strftime("%Y-%m-%d")
    out_spending = df_spending.copy()
    out_spending["지출날짜"] = out_spending["지출날짜"].dt.strftime("%Y-%m-%d")

    with pd.ExcelWriter(tagged_path, engine="openpyxl") as writer:
        out_all.to_excel(writer, index=False, sheet_name="전체(태깅포함)")
        out_spending.to_excel(writer, index=False, sheet_name="지출분석용")

    monthly = (
        df_spending.assign(연월=df_spending["지출날짜"].dt.to_period("M"))
        .groupby("연월", as_index=False)
        .agg(총_지출금액=("지출금액", "sum"), 거래건수=("지출금액", "size"))
        .sort_values("연월")
    )
    monthly["건당_평균지출"] = monthly["총_지출금액"] / monthly["거래건수"]
    monthly["연월"] = monthly["연월"].astype(str)
    with pd.ExcelWriter(monthly_path, engine="openpyxl") as writer:
        monthly.to_excel(writer, index=False, sheet_name="월별총지출")

    wb = load_workbook(monthly_path)
    ws = wb["월별총지출"]
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 16
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
        row[3].number_format = "#,##0"
    wb.save(monthly_path)

    logger("엑셀 출력 완료")
    return tagged_path, monthly_path


class SpendingTaggerApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("지출 태깅 & 차트 생성기 (macOS)")
        self.root.geometry("980x700")
        set_korean_font()

        self.files: list[Path] = []
        self.ontology: dict[str, list[str]] = load_ontology()
        self.mapping_profiles: list[dict] = load_mapping_profiles()
        if not ONTOLOGY_PATH.exists():
            save_ontology(self.ontology)
        if not MAPPING_PROFILES_PATH.exists():
            save_mapping_profiles(self.mapping_profiles)
        self._build_ui()

    def _build_ui(self) -> None:
        frame = ttk.Frame(self.root, padding=12)
        frame.pack(fill="both", expand=True)

        top = ttk.LabelFrame(frame, text="입력 파일")
        top.pack(fill="x")

        self.file_list = tk.Listbox(top, height=6)
        self.file_list.grid(row=0, column=0, rowspan=3, sticky="nsew", padx=(8, 8), pady=8)
        top.columnconfigure(0, weight=1)

        ttk.Button(top, text="파일 추가", command=self.add_files).grid(row=0, column=1, sticky="ew", padx=(0, 8), pady=(8, 4))
        ttk.Button(top, text="선택 제거", command=self.remove_selected_file).grid(
            row=1, column=1, sticky="ew", padx=(0, 8), pady=4
        )
        ttk.Button(top, text="모두 비우기", command=self.clear_files).grid(row=2, column=1, sticky="ew", padx=(0, 8), pady=4)

        out_frame = ttk.LabelFrame(frame, text="출력 위치")
        out_frame.pack(fill="x", pady=(10, 0))
        self.output_dir_var = tk.StringVar(value=str(Path.cwd() / "app_output"))
        ttk.Entry(out_frame, textvariable=self.output_dir_var).grid(row=0, column=0, sticky="ew", padx=8, pady=8)
        out_frame.columnconfigure(0, weight=1)
        ttk.Button(out_frame, text="폴더 선택", command=self.pick_output_dir).grid(row=0, column=1, padx=(0, 8), pady=8)

        ai_frame = ttk.LabelFrame(frame, text="자동 태깅 설정")
        ai_frame.pack(fill="x", pady=(10, 0))
        self.use_ai_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(ai_frame, text="GPT 자동 태깅 사용", variable=self.use_ai_var).grid(
            row=0, column=0, columnspan=2, sticky="w", padx=8, pady=(8, 4)
        )

        ttk.Label(ai_frame, text="OpenAI API Key").grid(row=1, column=0, sticky="w", padx=8)
        self.api_key_var = tk.StringVar(value=os.getenv("OPENAI_API_KEY", ""))
        self.api_key_entry = ttk.Entry(ai_frame, textvariable=self.api_key_var, show="*")
        self.api_key_entry.grid(row=1, column=1, sticky="ew", padx=(0, 8), pady=4)
        self._bind_paste_shortcuts(self.api_key_entry)
        ttk.Button(ai_frame, text="붙여넣기", command=self.paste_api_key_from_clipboard).grid(
            row=1, column=2, sticky="ew", padx=(0, 8), pady=4
        )

        ttk.Label(ai_frame, text="모델").grid(row=2, column=0, sticky="w", padx=8)
        self.model_var = tk.StringVar(value="gpt-4o-mini")
        ttk.Entry(ai_frame, textvariable=self.model_var).grid(row=2, column=1, sticky="ew", padx=(0, 8), pady=4)

        ttk.Label(ai_frame, text="자동확정 신뢰도(0~1)").grid(row=3, column=0, sticky="w", padx=8, pady=(0, 8))
        self.threshold_var = tk.StringVar(value="0.75")
        ttk.Entry(ai_frame, textvariable=self.threshold_var).grid(row=3, column=1, sticky="ew", padx=(0, 8), pady=(0, 8))
        ai_frame.columnconfigure(1, weight=1)

        ontology_frame = ttk.LabelFrame(frame, text="카테고리/세부카테고리 온톨로지")
        ontology_frame.pack(fill="x", pady=(10, 0))
        self.ontology_summary_var = tk.StringVar()
        ttk.Label(ontology_frame, textvariable=self.ontology_summary_var).grid(row=0, column=0, sticky="w", padx=8, pady=8)
        ttk.Button(ontology_frame, text="온톨로지 편집", command=self.open_ontology_editor).grid(
            row=0, column=1, padx=(0, 8), pady=8
        )
        ttk.Button(ontology_frame, text="기본값 복원", command=self.reset_ontology_default).grid(
            row=0, column=2, padx=(0, 8), pady=8
        )
        ontology_frame.columnconfigure(0, weight=1)

        map_frame = ttk.LabelFrame(frame, text="포맷 매핑 프로필")
        map_frame.pack(fill="x", pady=(10, 0))
        self.mapping_profile_summary_var = tk.StringVar()
        ttk.Label(map_frame, textvariable=self.mapping_profile_summary_var).grid(row=0, column=0, sticky="w", padx=8, pady=8)
        ttk.Button(map_frame, text="프로필 초기화", command=self.reset_mapping_profiles).grid(row=0, column=1, padx=(0, 8), pady=8)
        map_frame.columnconfigure(0, weight=1)

        run_frame = ttk.Frame(frame)
        run_frame.pack(fill="x", pady=(10, 0))
        self.run_btn = ttk.Button(run_frame, text="실행", command=self.run_pipeline)
        self.run_btn.pack(side="left")

        ttk.Label(
            run_frame,
            text='제외 카테고리: "제외", "금융/투자" (차트/지출집계에서 제외)',
        ).pack(side="left", padx=(12, 0))

        log_frame = ttk.LabelFrame(frame, text="로그")
        log_frame.pack(fill="both", expand=True, pady=(10, 0))
        self.log_text = tk.Text(log_frame, height=18, wrap="word")
        self.log_text.pack(fill="both", expand=True, padx=8, pady=8)
        self.refresh_ontology_summary()
        self.refresh_mapping_profile_summary()

    def log(self, msg: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert("end", f"[{timestamp}] {msg}\n")
        self.log_text.see("end")
        self.root.update_idletasks()

    def add_files(self) -> None:
        picked = filedialog.askopenfilenames(
            parent=self.root,
            title="거래내역 엑셀 선택",
            filetypes=[("Excel files", "*.xls *.xlsx"), ("All files", "*.*")],
        )
        for f in picked:
            p = Path(f)
            if p not in self.files:
                self.files.append(p)
                self.file_list.insert("end", p.as_posix())

    def remove_selected_file(self) -> None:
        sel = list(self.file_list.curselection())
        if not sel:
            return
        for i in reversed(sel):
            self.file_list.delete(i)
            del self.files[i]

    def clear_files(self) -> None:
        self.file_list.delete(0, "end")
        self.files.clear()

    def pick_output_dir(self) -> None:
        d = filedialog.askdirectory(parent=self.root, title="출력 폴더 선택")
        if d:
            self.output_dir_var.set(d)

    def _bind_paste_shortcuts(self, widget: ttk.Entry) -> None:
        widget.bind("<Command-v>", lambda e: self._paste_into_widget(widget))
        widget.bind("<Command-V>", lambda e: self._paste_into_widget(widget))
        widget.bind("<Control-v>", lambda e: self._paste_into_widget(widget))
        widget.bind("<Control-V>", lambda e: self._paste_into_widget(widget))
        widget.bind("<Shift-Insert>", lambda e: self._paste_into_widget(widget))

    def _paste_into_widget(self, widget: ttk.Entry) -> str:
        try:
            text = self.root.clipboard_get()
        except tk.TclError:
            return "break"
        try:
            widget.delete("sel.first", "sel.last")
        except tk.TclError:
            pass
        widget.insert("insert", text)
        return "break"

    def paste_api_key_from_clipboard(self) -> None:
        try:
            text = self.root.clipboard_get()
        except tk.TclError:
            messagebox.showwarning("확인", "클립보드에서 텍스트를 읽을 수 없습니다.")
            return
        self.api_key_var.set(str(text).strip())
        self.log("API Key를 클립보드에서 붙여넣었습니다.")

    def refresh_ontology_summary(self) -> None:
        categories = len(self.ontology)
        subcategories = sum(len(v) for v in self.ontology.values())
        self.ontology_summary_var.set(
            f"현재 온톨로지: 카테고리 {categories}개, 세부 카테고리 {subcategories}개 "
            f"(저장 파일: {ONTOLOGY_PATH.name})"
        )

    def refresh_mapping_profile_summary(self) -> None:
        self.mapping_profile_summary_var.set(
            f"현재 저장된 포맷 매핑 프로필: {len(self.mapping_profiles)}개 "
            f"(저장 파일: {MAPPING_PROFILES_PATH.name})"
        )

    def open_ontology_editor(self) -> None:
        dlg = OntologyEditorDialog(self.root, self.ontology)
        self.root.wait_window(dlg)
        if dlg.result is None:
            return
        self.ontology = normalize_ontology(dlg.result)
        save_ontology(self.ontology)
        self.refresh_ontology_summary()
        self.log("온톨로지 저장 완료")

    def reset_ontology_default(self) -> None:
        ok = messagebox.askyesno("확인", "온톨로지를 기본값으로 복원할까요?")
        if not ok:
            return
        self.ontology = get_default_ontology()
        save_ontology(self.ontology)
        self.refresh_ontology_summary()
        self.log("온톨로지를 기본값으로 복원")

    def reset_mapping_profiles(self) -> None:
        ok = messagebox.askyesno("확인", "저장된 포맷 매핑 프로필을 모두 초기화할까요?")
        if not ok:
            return
        self.mapping_profiles = []
        save_mapping_profiles(self.mapping_profiles)
        self.refresh_mapping_profile_summary()
        self.log("포맷 매핑 프로필 초기화 완료")

    def prompt_column_mapping(self, source_path: Path, raw_df: pd.DataFrame) -> dict | None:
        dlg = ColumnMappingDialog(self.root, source_path, raw_df)
        self.root.wait_window(dlg)
        result = dlg.result
        if result is None:
            return None
        if result.get("save_profile"):
            _upsert_mapping_profile(self.mapping_profiles, result)
            save_mapping_profiles(self.mapping_profiles)
            self.refresh_mapping_profile_summary()
            self.log(f"매핑 프로필 저장: {result.get('profile_name', '')}")
        return result

    def _review_uncertain(self, df: pd.DataFrame, indices: list[int], kb: KnowledgeBase) -> pd.DataFrame:
        if not indices:
            return df

        cancel_all = False
        reviewed = 0
        ontology_dirty = False
        for i, idx in enumerate(indices, start=1):
            if cancel_all:
                break

            row = df.loc[idx]
            suggestion = None
            rc = str(row.get("추천카테고리", "")).strip()
            rs = str(row.get("추천세부카테고리", "")).strip()
            conf = float(row.get("신뢰도", 0.0) or 0.0)
            source = str(row.get("자동태깅근거", "")).replace("review:", "")
            if rc:
                suggestion = TagSuggestion(rc, rs or "미분류", conf, source)

            self.log(f"검토 {i}/{len(indices)}: {row['사용처']} ({row['지출금액']:,.0f}원)")
            dlg = ReviewDialog(self.root, row, suggestion, self.ontology)
            self.root.wait_window(dlg)
            if dlg.ontology_changed:
                ontology_dirty = True

            if dlg.result == "__CANCEL_ALL__":
                cancel_all = True
                continue
            if dlg.result is None:
                continue

            cat, sub = dlg.result
            df.at[idx, "카테고리"] = cat
            df.at[idx, "세부 카테고리"] = sub or "미분류"
            df.at[idx, "자동태깅근거"] = "manual_review"
            df.at[idx, "신뢰도"] = 1.0
            reviewed += 1

        if ontology_dirty:
            save_ontology(self.ontology)
            self.refresh_ontology_summary()
            self.log("검토 팝업에서 수정한 온톨로지 저장 완료")

        self.log(f"검토 반영 완료: {reviewed}건")
        return df

    def run_pipeline(self) -> None:
        if not self.files:
            messagebox.showwarning("확인", "입력 파일을 1개 이상 선택해 주세요.")
            return

        try:
            threshold = float(self.threshold_var.get().strip())
        except Exception:
            messagebox.showerror("오류", "신뢰도 값은 0~1 숫자로 입력해 주세요.")
            return
        threshold = max(0.0, min(1.0, threshold))

        output_root = Path(self.output_dir_var.get().strip() or ".")
        run_dir = output_root / f"결과_{datetime.now():%Y%m%d_%H%M%S}"
        charts_dir = run_dir / "charts"

        self.run_btn.configure(state="disabled")
        self.log("작업 시작")
        try:
            run_dir.mkdir(parents=True, exist_ok=True)
            api_key = self.api_key_var.get().strip()
            model_name = self.model_var.get().strip() or "gpt-4o-mini"
            ai_mapper = None
            if api_key:
                ai_mapper = OpenAIColumnMapper(api_key=api_key, model=model_name)
                self.log("GPT 포맷 파싱 폴백 활성화")
            else:
                self.log("GPT 포맷 파싱 폴백 비활성화(API Key 없음)")

            df = load_and_merge(
                self.files,
                self.log,
                mapping_profiles=self.mapping_profiles,
                mapping_prompt=self.prompt_column_mapping,
                ai_mapping_suggester=ai_mapper,
            )
            save_mapping_profiles(self.mapping_profiles)
            self.refresh_mapping_profile_summary()
            self.log(f"입력 통합: {len(df)}건")

            if Path("통합지출내역.xlsx").exists():
                try:
                    ref = pd.read_excel("통합지출내역.xlsx")
                    ref = normalize_columns(ref, "통합지출내역.xlsx")
                    combined_for_kb = pd.concat([df, ref], ignore_index=True)
                    self.log(f"참조 라벨 데이터 사용: {len(ref)}건")
                except Exception:
                    combined_for_kb = df
            else:
                combined_for_kb = df

            kb = build_knowledge(combined_for_kb, self.ontology, self.log)

            ai_tagger = None
            if self.use_ai_var.get():
                if not api_key:
                    raise ValueError("GPT 자동 태깅을 켰다면 API Key를 입력해야 합니다.")
                ai_tagger = OpenAITagger(api_key=api_key, model=model_name)
                self.log("GPT 자동 태깅 활성화")
            else:
                self.log("GPT 자동 태깅 비활성화(규칙 기반 + 수동검토)")

            tagged, uncertain = auto_tag(df, kb, ai_tagger, threshold, self.log)

            if uncertain:
                mandatory_indices = [
                    idx
                    for idx in uncertain
                    if "ambiguous_manual_required" in str(tagged.at[idx, "자동태깅근거"])
                ]
                optional_indices = [idx for idx in uncertain if idx not in mandatory_indices]

                if mandatory_indices:
                    messagebox.showinfo(
                        "필수 분류 필요",
                        f'미분류 또는 "타행송금/타행이체" 성격의 항목 {len(mandatory_indices)}건은 '
                        "반드시 사용 목적을 분류해야 합니다.",
                    )
                    tagged = self._review_uncertain(tagged, uncertain, kb)

                    unresolved = [
                        idx
                        for idx in mandatory_indices
                        if "ambiguous_manual_required" in str(tagged.at[idx, "자동태깅근거"])
                    ]
                    if unresolved:
                        raise ValueError(
                            f'미분류/타행송금 계열 항목 {len(unresolved)}건이 아직 분류되지 않았습니다. '
                            "해당 항목을 분류 후 다시 실행해 주세요."
                        )
                elif optional_indices:
                    do_review = messagebox.askyesno(
                        "검토 필요",
                        f"불확실 항목 {len(optional_indices)}건이 있습니다.\n지금 수동 검토를 진행할까요?",
                    )
                    if do_review:
                        tagged = self._review_uncertain(tagged, optional_indices, kb)
                    else:
                        self.log("수동 검토를 건너뜀")

            spend_df = spending_only(tagged)
            self.log(f"지출 분석 대상: {len(spend_df)}건")
            if spend_df.empty:
                raise ValueError("지출 분석 대상이 없습니다. 카테고리 태깅 또는 제외 규칙을 확인해 주세요.")

            tagged_path, monthly_path = export_excels(tagged, spend_df, run_dir, self.log)
            chart_paths = export_charts(spend_df, charts_dir, self.log)
            update_memory(tagged)
            self.log(f"태깅 메모리 저장: {MEMORY_PATH.as_posix()}")

            summary = (
                f"완료\n\n"
                f"- 결과 폴더: {run_dir.as_posix()}\n"
                f"- 태깅 엑셀: {tagged_path.name}\n"
                f"- 월별 요약 엑셀: {monthly_path.name}\n"
                f"- 차트 수: {len(chart_paths)}개"
            )
            messagebox.showinfo("완료", summary)
            self.log("작업 완료")
        except Exception as e:
            self.log(f"오류: {type(e).__name__}: {e}")
            messagebox.showerror("오류", f"{type(e).__name__}: {e}")
        finally:
            self.run_btn.configure(state="normal")


def main() -> None:
    root = tk.Tk()
    app = SpendingTaggerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
